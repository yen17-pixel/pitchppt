import io
import json
import streamlit as st
from pptx import Presentation
from google import genai

# Configure Page Layout
st.set_page_config(page_title="AI Proposal Generator", page_icon="📊", layout="wide")

st.title("📊 Technical Proposal & Pitch Deck Generator")
st.write("Transform raw client requirements and meeting notes into a structured PowerPoint proposal.")

# Sidebar - Configuration & Case Studies
with st.sidebar:
    st.header("⚙️ Settings & Credentials")
    api_key = st.text_input("Enter Gemini API Key", type="password")
    
    st.markdown("---")
    st.subheader("📁 Embedded Case Studies")
    st.caption("The AI agent queries these past wins to auto-insert proof points:")
    
    # Pre-loaded sample case study repository
    sample_case_studies = [
        {
            "service": "Carbon Footprint & Offset Assessment",
            "challenge": "Client required Scope 1 & 2 baseline emissions data for multi-facility operations.",
            "solution": "Executed energy audits, modeled carbon baselines, and established a 3-year reduction roadmap.",
            "metrics": ["30% emissions reduction target set", "Full compliance achieved in 4 weeks", "Identified $50k annual savings"]
        },
        {
            "service": "ESG & HSE Compliance Audit",
            "challenge": "Needed rapid gap analysis for ESG reporting and occupational health and safety compliance.",
            "solution": "Conducted multi-site inspections, created an ESG framework, and provided compliance training.",
            "metrics": ["Zero critical non-conformances", "Framework delivered in 3 weeks", "100% staff compliance"]
        }
    ]
    st.json(sample_case_studies)

# Main Form Layout
col1, col2 = st.columns([1, 1])

with col1:
    client_name = st.text_input("Client / Company Name", placeholder="e.g. Pharmatec Industries")
    service_type = st.selectbox(
        "Primary Service Line",
        ["Carbon Footprint & ESG Audit", "HSE & Environmental Compliance", "Sustainability Advisory", "Custom Technical Solution"]
    )

with col2:
    project_title = st.text_input("Proposal Presentation Title", placeholder="e.g. Technical Proposal for Carbon Assessment")
    include_case_study = st.checkbox("Automatically Include Relevant Case Study Slide", value=True)

client_brief = st.text_area(
    "Paste Client Meeting Notes / RFP Requirements",
    height=180,
    placeholder="Paste raw notes here... e.g. Client needs an energy audit across 3 manufacturing sites within 6 weeks to achieve compliance..."
)

# Action Trigger
if st.button("🚀 Generate Client Presentation", type="primary"):
    if not api_key:
        st.error("Please enter your Gemini API Key in the sidebar.")
    elif not client_brief or not client_name:
        st.warning("Please provide both Client Name and Client Requirements.")
    else:
        with st.spinner("AI Brain is analyzing brief, selecting case studies, and creating slides..."):
            try:
                # 1. Initialize Gemini Client
                client = genai.Client(api_key=api_key)

                # 2. Build Structured Reasoning Prompt
                prompt = f"""
                You are a Lead Solution Architect. Analyze the client details below and create a structured JSON for a high-converting technical proposal deck.

                Client Name: {client_name}
                Service Line: {service_type}
                Client Notes:
                {client_brief}

                Available Case Studies:
                {json.dumps(sample_case_studies)}

                Return ONLY a valid JSON object with no extra formatting matching this exact structure:
                {{
                    "presentation_title": "{project_title or 'Technical & Strategic Proposal'}",
                    "slides": [
                        {{
                            "slide_title": "Executive Summary & Objectives",
                            "points": ["3 key bullet points on client challenges, urgency, and proposed strategic solution"]
                        }},
                        {{
                            "slide_title": "Technical Scope of Work",
                            "points": ["3-4 bullet points outlining project deliverables, methodology, and scope"]
                        }},
                        {{
                            "slide_title": "Execution Roadmap & Milestones",
                            "points": ["3 bullet points breaking down project phases and timeline"]
                        }}
                    ],
                    "matched_case_study": {{
                        "title": "Proof of Performance: Case Study",
                        "challenge": "Summary of similar past challenge solved",
                        "solution": "How we solved it",
                        "metrics": ["3 short metric results with numbers"]
                    }}
                }}
                """

                response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=prompt,
                )

                # Clean and parse JSON response
                raw_json = response.text.replace("```json", "").replace("```", "").strip()
                proposal_data = json.loads(raw_json)

                # 3. Build PowerPoint File using python-pptx
                prs = Presentation()
                bullet_layout = prs.slide_layouts[1]

                # Title Slide
                title_slide = prs.slides.add_slide(prs.slide_layouts[0])
                title_slide.shapes.title.text = proposal_data.get("presentation_title", "Technical Proposal")
                title_slide.placeholders[1].text = f"Prepared for: {client_name}\nService: {service_type}"

                # Content Slides
                for slide_info in proposal_data.get("slides", []):
                    slide = prs.slides.add_slide(bullet_layout)
                    slide.shapes.title.text = slide_info["slide_title"]
                    tf = slide.shapes.placeholders[1].text_frame
                    
                    points = slide_info.get("points", [])
                    if points:
                        tf.text = points[0]
                        for p in points[1:]:
                            para = tf.add_paragraph()
                            para.text = p

                # Case Study Slide
                if include_case_study and "matched_case_study" in proposal_data:
                    cs = proposal_data["matched_case_study"]
                    cs_slide = prs.slides.add_slide(bullet_layout)
                    cs_slide.shapes.title.text = cs.get("title", "Case Study")
                    
                    tf = cs_slide.shapes.placeholders[1].text_frame
                    tf.text = f"Challenge: {cs.get('challenge', '')}"
                    
                    p_sol = tf.add_paragraph()
                    p_sol.text = f"Solution: {cs.get('solution', '')}"
                    
                    p_res = tf.add_paragraph()
                    p_res.text = "Key Results & ROI:"
                    for m in cs.get("metrics", []):
                        p_m = tf.add_paragraph()
                        p_m.text = f"• {m}"
                        p_m.level = 1

                # 4. Streamlit Download Handler (Save to In-Memory Buffer)
                ppt_buffer = io.BytesIO()
                prs.save(ppt_buffer)
                ppt_buffer.seek(0)

                st.success("🎉 Proposal presentation successfully generated!")
                
                # Show structured preview in UI
                with st.expander("Preview Generated Structure"):
                    st.json(proposal_data)

                # Download Button
                st.download_button(
                    label="📥 Download PowerPoint (.pptx)",
                    data=ppt_buffer,
                    file_name=f"{client_name.replace(' ', '_')}_Proposal.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    type="primary"
                )

            except Exception as e:
                st.error(f"Execution Error: {str(e)}")